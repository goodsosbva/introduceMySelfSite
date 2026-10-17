from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).parent
FONT = "Malgun Gothic"
NAVY = "17324D"
SLATE = "425466"
PALE = "EEF3F7"
LINE = "D9E1E8"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, data in kwargs.items():
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in data.items():
            element.set(qn("w:" + key), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_run_font(run, size=10.5, bold=False, color="000000"):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def clear_paragraph(p):
    p._element.clear_content()


def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    pf = normal.paragraph_format
    pf.space_after = Pt(5)
    pf.line_spacing = 1.28

    for name, size in (("Title", 25), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = ""


def paragraph(doc, text="", style=None, before=0, after=5, bold_prefix=None, size=10.5, color="000000"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.28
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, size=size, bold=True, color=color)
        run = p.add_run(text[len(bold_prefix) :])
        set_run_font(run, size=size, color=color)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color)
    return p


def heading(doc, text, level=1, before=18, after=8):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 11.5, bold=True, color="000000")
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="Normal")
    p.style = doc.styles["Normal"]
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.35)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.24
    bullet_run = p.add_run("• ")
    set_run_font(bullet_run, size=10, color=NAVY)
    run = p.add_run(text)
    set_run_font(run, size=10)
    return p


def contact_block(doc, role):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(role)
    set_run_font(run, size=12, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    for index, item in enumerate((
        "010-3223-8549",
        "goodsosbva@naver.com",
        "github.com/goodsosbva",
        "goodsosbva.tistory.com",
    )):
        if index:
            separator = p.add_run("  |  ")
            set_run_font(separator, size=9.5, color=SLATE)
        run = p.add_run(item)
        set_run_font(run, size=9.5, color=SLATE)


def two_column_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(3.25)
    table.columns[1].width = Cm(13.0)
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(3.25)
        cells[1].width = Cm(13.0)
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, top={"val": "single", "sz": 4, "color": LINE}, bottom={"val": "single", "sz": 4, "color": LINE}, left={"val": "single", "sz": 4, "color": LINE}, right={"val": "single", "sz": 4, "color": LINE})
        set_cell_shading(cells[0], PALE)
        for cell, text, bold, color in ((cells[0], left, True, NAVY), (cells[1], right, False, "000000")):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(text)
            set_run_font(run, size=9.5, bold=bold, color=color)
    return table


def start_doc(title, role):
    doc = Document()
    configure(doc)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    set_run_font(run, size=25, bold=True)
    contact_block(doc, role)
    return doc


def build_resume():
    doc = start_doc("권현성 이력서", "Frontend Developer")
    heading(doc, "프로필", before=0)
    paragraph(doc, "Vue와 TypeScript를 기반으로 클라우드 프로비저닝과 모니터링 운영 화면을 개발해 온 프론트엔드 개발자입니다. 화면을 구현하기 전에 데이터 흐름과 실행 조건을 먼저 정리하고, 로딩·빈 결과·권한 부족·오류까지 사용자가 다음 행동을 판단할 수 있는 상태로 설계합니다. 공통 UI와 자동화 검증을 통해 팀이 계속 개선할 수 있는 코드를 만드는 데 집중합니다.", after=8)

    heading(doc, "경력")
    paragraph(doc, "오케스트로 주식회사  2023 09 현재", bold_prefix="오케스트로 주식회사", after=2)
    paragraph(doc, "Frontend Developer  클라우드 관리 플랫폼 CMP", size=9.5, color=SLATE, after=4)
    bullet(doc, "Kubernetes, OpenShift, AWS의 자원 모델과 입력 규칙을 공통 폼 구조로 정리해 하나의 서비스 카탈로그 흐름에서 프로비저닝을 처리했습니다.")
    bullet(doc, "권한, 자원 상태, 입력 완료 여부를 실행 조건과 연결하고 오류 피드백을 구성해 잘못된 API 요청을 화면 단계에서 차단했습니다.")
    bullet(doc, "Alarm Rule, Alarm Detail, Alarm History를 연결하고 24시간·7일 기준의 통계와 분포를 시각화해 운영자가 상태와 원인을 탐색하도록 했습니다.")
    bullet(doc, "공통 테이블의 말줄임과 컬럼 리사이즈 문제를 실제 렌더링 너비 측정 방식으로 개선했습니다.")
    bullet(doc, "Playwright 기반 Common Tester를 구성해 문서 기반 검증 규칙을 실제 브라우저 시나리오로 실행하고 실패 원인을 추적했습니다.")

    paragraph(doc, "주식회사 카랑  2023 02 2023 07", bold_prefix="주식회사 카랑", before=8, after=2)
    paragraph(doc, "Frontend Developer  자동차 수리 예약 관리", size=9.5, color=SLATE, after=4)
    bullet(doc, "정비사가 예약 완료 건의 차량 정보, 고객 요청, 일정을 조회하고 작업 상태를 관리하는 웹 화면을 구현했습니다.")
    bullet(doc, "REST API 요청과 응답 상태를 연결하고 반응형 UI와 반복 요소의 컴포넌트화를 적용했습니다.")

    paragraph(doc, "주식회사 뉴스젤리  2022 12 2023 01", bold_prefix="주식회사 뉴스젤리", before=8, after=2)
    paragraph(doc, "Frontend Developer Intern  데이터 시각화 서비스", size=9.5, color=SLATE, after=4)
    bullet(doc, "Vanilla JavaScript 기반 화면을 Vue 구조로 마이그레이션하고 중복 UI와 로직을 재사용 가능한 컴포넌트로 분리했습니다.")

    heading(doc, "기술")
    two_column_table(doc, [
        ("언어와 UI", "JavaScript ES6  TypeScript  Vue 3  React  Next.js  HTML5  CSS3  Sass"),
        ("상태와 데이터", "Pinia  TanStack Vue Query  Vee Validate  Yup  REST API  SQL  MySQL"),
        ("도구와 품질", "Vite  Webpack Module Federation  Playwright  Storybook  Jenkins  GitHub Actions"),
        ("플랫폼", "Kubernetes  OpenShift  AWS  Docker  Nginx  Linux  Node.js  Express  NestJS  FastAPI"),
    ])

    heading(doc, "대표 프로젝트")
    paragraph(doc, "CityWatch FE Lab  개인 프로젝트", bold_prefix="CityWatch FE Lab", after=2)
    paragraph(doc, "Next.js Host, Vite Analytics Remote, 실시간 Node 서버와 공유 UI·타입 패키지를 npm Workspaces 모노레포로 구성했습니다. Module Federation 원격 로딩, WebSocket 장애 시 Polling 전환, 런타임 데이터 검증, 오류 상태와 재시도를 구현했습니다.", size=10, after=4)
    paragraph(doc, "GitHub  github.com/goodsosbva/citywatch-fe-lab", size=9.5, color=SLATE, after=4)

    heading(doc, "학력과 자격")
    two_column_table(doc, [
        ("학력", "강원대학교 컴퓨터공학과 졸업  2019 03 2021 02\n김포대학교 컴퓨터네트워크과 졸업  2017 03 2019 02"),
        ("자격", "Certified Kubernetes Administrator CKA  2026 01\nSQLD  2025 04  리눅스마스터 2급  2023 06  정보처리기사  2021 08"),
        ("기여", "전자정부 표준프레임워크 다국어 기능 개선 제안 및 기여  감사장 수상  2025 10"),
    ])
    doc.save(OUT / "권현성_이력서.docx")


def build_career():
    doc = start_doc("권현성 경력기술서", "Frontend Developer")
    heading(doc, "경력 요약", before=0)
    paragraph(doc, "클라우드 관리 플랫폼에서 프로비저닝, 모니터링, 알람, 공통 UI를 개발했습니다. 서로 다른 자원 모델과 운영 조건을 사용자가 판단 가능한 화면 흐름으로 정리하고, 입력 검증과 오류 피드백으로 잘못된 실행을 줄이는 데 집중했습니다. Vue 3, TypeScript, Pinia, Vite, Ant Design Vue 환경에서 REST API 연동과 Jenkins 기반 배포, Jira·Bitbucket 협업을 경험했습니다.", after=10)

    heading(doc, "오케스트로 주식회사 2023 09 현재")
    paragraph(doc, "Frontend Developer  클라우드 관리 플랫폼 CMP", size=10, color=SLATE, after=7)
    heading(doc, "프로비저닝 동적 폼", level=2, before=6)
    paragraph(doc, "Kubernetes, OpenShift, AWS의 자원 모델과 입력 규칙을 플랫폼별 분기와 공통 폼 구조로 정리했습니다. 서비스 카탈로그에서 자원 생성까지 이어지는 흐름을 구성하면서 플랫폼마다 달라지는 필드와 검증 규칙은 동적 컴포넌트로 분리했습니다.")
    bullet(doc, "기관별 권한, 자원 상태, 입력 완료 여부를 실행 조건에 연결했습니다.")
    bullet(doc, "실행 가능 여부와 오류 사유를 화면에서 확인할 수 있도록 피드백을 구성했습니다.")
    bullet(doc, "잘못된 요청이 API 호출로 이어지지 않도록 입력과 상태를 사전에 검증했습니다.")

    heading(doc, "모니터링과 알람", level=2, before=10)
    paragraph(doc, "Alarm Rule의 목록·상세·생성 화면과 Alarm Detail·History를 연결해 운영자가 자원별 알람 현황과 이력을 같은 탐색 흐름에서 확인하도록 구현했습니다. 서로 다른 API를 사용하는 통계 카드, 차트, 목록, 상세 패널이 일관된 조회 기준을 바라보도록 상태 흐름을 설계했습니다.")
    bullet(doc, "최근 24시간과 7일 데이터를 최다·평균·조용한 시간대·요일별 분포로 가공해 시각화했습니다.")
    bullet(doc, "페이지네이션과 재조회 시점을 맞추고 권한 부족, 처리 중, 빈 결과, 오류 상태를 구분했습니다.")
    bullet(doc, "플랫폼 null, Severity, 타임스탬프, 다국어 표기 규칙을 화면 전반에서 통일했습니다.")

    heading(doc, "공통 테이블 렌더링 문제 개선", level=2, before=10)
    paragraph(doc, "공통 테이블에서 긴 텍스트가 말줄임 처리되지 않고 리사이즈를 반복할수록 전체 너비가 넓어지는 문제를 해결했습니다. 발생 일시 컬럼에 명시적 너비가 없어 셀이 내용만큼 확장되는 원인을 먼저 해결한 뒤, 상태에 저장된 컬럼 너비와 브라우저가 실제 렌더링한 헤더 너비의 차이가 반복 계산에서 오차를 누적한다는 점을 확인했습니다.")
    paragraph(doc, "모든 컬럼에 고정 너비와 table layout fixed를 적용하는 방법은 간단했지만 공통 테이블의 유연성과 사용자의 컬럼 조절 기능을 제한했습니다. 선언된 너비만 기준으로 계산하는 방법도 실제 렌더링 결과와의 차이를 해결하지 못했습니다. 최종적으로 렌더링 후 각 헤더 셀의 실제 너비를 측정해 컬럼 상태에 반영했고, 조건부 컬럼 유무와 관계없이 반복 리사이즈와 말줄임이 안정적으로 동작하도록 했습니다.")

    heading(doc, "검증 자동화", level=2, before=10)
    paragraph(doc, "문서로 관리되던 테스트 규칙을 Playwright 시나리오로 변환해 실제 브라우저에서 실행하는 Common Tester를 구성했습니다. 성공 경로뿐 아니라 중복 요청, 응답 지연, 승인 실패를 실행 가능한 테스트로 남겼으며, 실패 지점을 문서 해석·실행 환경·화면 동작으로 구분해 확인 범위를 빠르게 좁히도록 했습니다.")

    heading(doc, "주식회사 카랑 2023 02 2023 07")
    paragraph(doc, "Frontend Developer  자동차 수리 예약 관리", size=10, color=SLATE, after=5)
    bullet(doc, "정비사가 예약 완료 건의 차량 정보, 고객 요청, 일정을 조회하고 작업 상태를 관리하는 화면을 개발했습니다.")
    bullet(doc, "사용자 예약 단계와 운영자 관리 화면의 UI 및 상태 흐름을 구현했습니다.")
    bullet(doc, "REST API 연동, 로딩·오류 상태 처리, 반응형 UI, 반복 요소 컴포넌트화를 적용했습니다.")

    heading(doc, "주식회사 뉴스젤리 2022 12 2023 01")
    paragraph(doc, "Frontend Developer Intern  데이터 시각화 서비스", size=10, color=SLATE, after=5)
    bullet(doc, "Vanilla JavaScript 기반 서비스를 Vue 기반 구조로 마이그레이션했습니다.")
    bullet(doc, "기존 동작을 유지하면서 중복 UI와 로직을 재사용 가능한 컴포넌트 구조로 분리했습니다.")

    heading(doc, "개인 프로젝트")
    paragraph(doc, "CityWatch FE Lab  2026 07 2026 08", bold_prefix="CityWatch FE Lab", after=3)
    bullet(doc, "Next.js Host, Vite Analytics Remote, 실시간 Node 서버를 독립 실행 영역으로 분리했습니다.")
    bullet(doc, "공유 UI와 API 타입 패키지를 포함한 npm Workspaces 모노레포를 구성했습니다.")
    bullet(doc, "Module Federation 원격 로딩, WebSocket 장애 시 Polling 전환, 런타임 데이터 검증, 오류 상태와 재시도를 구현했습니다.")
    bullet(doc, "Storybook, 테스트, 빌드 검증으로 화면과 문서뿐 아니라 실행 결과를 확인했습니다.")
    paragraph(doc, "GitHub  github.com/goodsosbva/citywatch-fe-lab", size=9.5, color=SLATE)
    doc.save(OUT / "권현성_경력기술서.docx")


def build_intro():
    doc = start_doc("권현성 자기소개서", "Frontend Developer")
    heading(doc, "문제를 끝까지 이해하는 개발자", before=0)
    paragraph(doc, "개발은 좋아하는 일을 오래 깊게 파고들 수 있는 방식이었습니다. 문제의 원인을 찾고 코드로 직접 해결해 누군가의 불편을 줄이는 과정에서 가장 큰 즐거움을 느꼈습니다. 지금도 화면을 먼저 만들기보다 데이터가 어디에서 시작해 어떤 조건을 지나 결과로 이어지는지 정리하고, 실패 상황을 재현해 원인을 좁힌 뒤 실행 가능한 규칙과 코드로 옮깁니다.")
    paragraph(doc, "업무 중 막혔던 지점과 해결 과정을 기록하는 습관도 이 태도에서 시작했습니다. 기록은 다음 문제를 더 빠르게 해결하는 자산이자, 동료에게 판단 근거를 분명히 설명하는 도구가 됩니다. 맡은 문제를 끝까지 이해하고 재현 가능한 형태로 남기는 개발자가 되고자 합니다.")

    heading(doc, "복잡한 조건을 사용자가 판단 가능한 화면으로")
    paragraph(doc, "클라우드 관리 플랫폼에서는 서버, 스토리지, 네트워크 자원의 생성과 조회에 기관별 권한, 알람 정책, 처리 상태가 함께 얽혀 있습니다. 내부 조건이 복잡할수록 사용자는 현재 상태와 다음 행동을 빠르게 판단할 수 있어야 한다고 생각합니다.")
    paragraph(doc, "오케스트로에서 Kubernetes, OpenShift, AWS 프로비저닝을 하나의 서비스 카탈로그 흐름으로 구성했습니다. 서로 다른 자원 모델과 입력 규칙은 공통 폼 구조와 동적 컴포넌트로 정리하고, 권한·자원 상태·입력 완료 여부를 실행 조건으로 연결했습니다. 또한 통계 카드, 차트, 목록, 상세 패널이 서로 다른 API를 사용해도 같은 조회 기준을 유지하도록 상태 흐름을 설계했습니다.")
    paragraph(doc, "페이지네이션과 재조회 시점을 맞추고 권한 부족, 처리 중, 빈 결과, 오류 상태를 명확히 구분했습니다. 반복되는 컬럼 제어, 상태 태그, 동적 폼은 공통 컴포넌트로 묶어 요구사항이 바뀌어도 화면마다 다른 기준이 생기지 않도록 했습니다.")

    heading(doc, "실제 렌더링 결과를 기준으로 문제를 해결한 경험")
    paragraph(doc, "가장 기억에 남는 기술적 문제는 공통 테이블에서 컬럼 너비를 줄여도 긴 텍스트가 말줄임 처리되지 않고, 리사이즈를 반복할수록 테이블이 넓어지는 오류였습니다. 먼저 발생 일시 컬럼에 명시적인 너비가 없어 셀이 텍스트 길이만큼 확장되는 문제를 해결했습니다. 그러나 리사이즈를 반복하면 상태에 저장된 너비와 브라우저가 실제로 렌더링한 헤더 너비의 차이가 누적되는 문제는 남아 있었습니다.")
    paragraph(doc, "모든 컬럼에 고정 너비와 table layout fixed를 적용하는 방법은 구현이 간단했지만 공통 테이블의 유연성과 사용자의 조절 기능을 제한했습니다. 선언된 컬럼 너비만 기준으로 계산하는 방법은 실제 렌더링 결과와의 차이를 해결하지 못했습니다. 최종적으로 렌더링 이후 각 헤더 셀의 실제 너비를 측정하고 이를 컬럼 상태에 반영했습니다. 사용자가 실제로 보고 있는 결과를 계산 기준으로 삼아야 한다고 판단했기 때문입니다. 이 경험을 통해 선언된 값뿐 아니라 브라우저가 계산한 결과까지 확인해야 프론트엔드 문제를 정확히 해결할 수 있음을 배웠습니다.")

    heading(doc, "자동화된 검증으로 신뢰를 만드는 방식")
    paragraph(doc, "AI가 생성한 코드의 가치는 빠른 생성 자체보다 실제 서비스 환경에서 의도대로 동작하는지 검증할 때 만들어진다고 생각합니다. 반복되는 화면 검증을 효율화하기 위해 문서 기반 테스트 규칙을 Playwright 시나리오로 변환하는 Common Tester를 구성했습니다.")
    paragraph(doc, "성공 시나리오만 확인하지 않고 중복 요청, 응답 지연, 승인 실패 같은 예외 조건도 실행 가능한 테스트로 남겼습니다. 테스트 실패 시에는 문서 해석, 실행 환경, 화면 동작 중 어느 단계에서 문제가 발생했는지 구분해 추적하도록 구성했습니다. 요구사항이 바뀌더라도 위험을 먼저 발견하고 확인 범위를 좁힐 수 있는 개발 흐름을 지향합니다.")

    heading(doc, "문제를 발견하면 행동으로 옮기는 태도")
    paragraph(doc, "전자정부 표준프레임워크의 다국어 기능을 살피던 중 보완할 수 있는 부분을 발견했습니다. 필요한 기능이 제공되지 않는다는 점을 확인한 뒤 개선안을 제안하고 직접 기여했으며, 2025년 감사장을 받았습니다. 작은 아이디어라도 실제 개선으로 연결하려면 문제를 발견하는 데서 멈추지 않고 근거를 갖춰 행동해야 한다는 점을 배웠습니다.")
    paragraph(doc, "한편 기술 전환 논의에서 갈등을 피하려고 의견을 충분히 내지 못했던 경험도 있습니다. 이후에는 기술 선택의 장단점과 위험을 문서로 정리하고, 가능하면 작은 POC로 빌드 결과와 개발 경험을 비교하려고 합니다. 의견이 채택되지 않더라도 근거를 분명히 전달하고, 결정된 방향에는 책임 있게 참여하겠습니다.")

    heading(doc, "앞으로의 방향")
    paragraph(doc, "화면을 만드는 데서 멈추지 않고 사용자가 신뢰할 수 있고 동료가 계속 개선할 수 있는 프론트엔드 시스템을 만들고 싶습니다. 복잡한 조건은 내부 구조로 감추되 사용자는 빠르게 판단할 수 있게 하고, 자동화된 테스트와 기록으로 실제 환경에서 안정적으로 동작하는지 끝까지 확인하겠습니다.")
    doc.save(OUT / "권현성_자기소개서.docx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_resume()
    build_career()
    build_intro()
