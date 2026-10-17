from pathlib import Path
from zipfile import ZipFile

from docx import Document


base = Path(__file__).parent / "final"
expected = {
    "권현성_이력서.docx": ["권현성 이력서", "오케스트로 주식회사", "CityWatch FE Lab"],
    "권현성_경력기술서.docx": ["권현성 경력기술서", "공통 테이블 렌더링 문제 개선", "CityWatch FE Lab"],
    "권현성_자기소개서.docx": ["권현성 자기소개서", "실제 렌더링 결과를 기준으로 문제를 해결한 경험", "앞으로의 방향"],
}

for name, terms in expected.items():
    path = base / name
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert all(term in text for term in terms), name
    assert doc.core_properties.author == "", name
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        assert "rsid" not in xml, name
    print(f"OK {name}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")

html = (Path(__file__).parent.parent / "portfolio" / "index.html").read_text(encoding="utf-8")
for token in ("lang=\"ko\"", "본문으로 바로가기", "<main id=\"main\">", "id=\"work\"", "id=\"career\"", "id=\"skills\"", "id=\"contact\"", "prefers-reduced-motion"):
    assert token in html, token
assert 'target="_blank" rel="noopener noreferrer"' in html
print("OK portfolio structure and accessibility checks")
